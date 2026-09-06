import os
from pptx import Presentation
from pptx.util import Inches, Pt

from tools.analytics import daily_close_report
from tools.alerts import get_reorder_suggestions


def create_business_report():
    os.makedirs("artifacts/reports", exist_ok=True)

    report = daily_close_report()
    reorder = get_reorder_suggestions()

    date = report["date"]

    file_path = f"artifacts/reports/business_report_{date}.pptx"

    presentation = Presentation()

    # Slide 1 - Title
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[0]
    )

    slide.shapes.title.text = "KiranaOS AI"
    slide.placeholders[1].text = (
        f"Daily Business Analysis Report\n{date}"
    )

    # Slide 2 - Sales Summary
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )

    slide.shapes.title.text = "Daily Sales Summary"

    textbox = slide.shapes.add_textbox(
        Inches(1),
        Inches(1.5),
        Inches(8),
        Inches(4)
    )

    text_frame = textbox.text_frame

    lines = [
        f"Number of Bills: {report['number_of_sales']}",
        f"Quantity Sold: {report['quantity_sold']}",
        f"Subtotal: ₹{report['subtotal']:.2f}",
        f"GST Collected: ₹{report['gst_collected']:.2f}",
        f"Total Sales: ₹{report['total_sales']:.2f}"
    ]

    for line in lines:
        paragraph = text_frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(22)

    # Slide 3 - Top Selling Products
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )

    slide.shapes.title.text = "Top Selling Products"

    textbox = slide.shapes.add_textbox(
        Inches(1),
        Inches(1.5),
        Inches(8),
        Inches(4.5)
    )

    text_frame = textbox.text_frame

    top_products = report["top_selling_products"]

    if top_products:
        for item in top_products:
            paragraph = text_frame.add_paragraph()
            paragraph.text = (
                f"{item['product']} - "
                f"{item['quantity_sold']:g} units - "
                f"₹{item['sales_value']:.2f}"
            )
            paragraph.font.size = Pt(20)
    else:
        paragraph = text_frame.add_paragraph()
        paragraph.text = "No sales recorded today."
        paragraph.font.size = Pt(22)

    # Slide 4 - Reorder Suggestions
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )

    slide.shapes.title.text = "Inventory Reorder Suggestions"

    textbox = slide.shapes.add_textbox(
        Inches(0.7),
        Inches(1.4),
        Inches(8.5),
        Inches(5)
    )

    text_frame = textbox.text_frame

    products = reorder["products"]

    if products:
        for item in products:
            paragraph = text_frame.add_paragraph()
            paragraph.text = (
                f"{item['product']} | "
                f"Stock: {item['current_stock']:g} {item['unit']} | "
                f"Reorder: {item['reorder_quantity']:g} | "
                f"Priority: {item['priority']}"
            )
            paragraph.font.size = Pt(17)
    else:
        paragraph = text_frame.add_paragraph()
        paragraph.text = "No products currently require reordering."
        paragraph.font.size = Pt(22)

    # Slide 5 - Business Insights
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )

    slide.shapes.title.text = "Business Insights"

    textbox = slide.shapes.add_textbox(
        Inches(0.8),
        Inches(1.4),
        Inches(8.5),
        Inches(5)
    )

    text_frame = textbox.text_frame

    if report["number_of_sales"] == 0:
        insights = [
            "No sales have been recorded today.",
            "Continue monitoring sales throughout the day.",
            "Review inventory before the next business cycle."
        ]
    else:
        insights = [
            f"{report['number_of_sales']} bills were completed today.",
            f"Total revenue generated: ₹{report['total_sales']:.2f}.",
            f"GST collected: ₹{report['gst_collected']:.2f}.",
            "Review reorder suggestions to maintain inventory availability."
        ]

    for insight in insights:
        paragraph = text_frame.add_paragraph()
        paragraph.text = "• " + insight
        paragraph.font.size = Pt(20)

    presentation.save(file_path)

    return {
        "success": True,
        "message": "Business analysis report created successfully.",
        "file_path": file_path
    }